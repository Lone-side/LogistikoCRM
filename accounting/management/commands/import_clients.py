import openpyxl
from django.core.management.base import BaseCommand
from django.db import transaction
from accounting.models import ClientProfile
from datetime import datetime
import sys


class Command(BaseCommand):
    help = 'Εισαγωγή πελατών από Excel με ΟΛΑ τα πεδία και advanced features'

    def add_arguments(self, parser):
        parser.add_argument(
            'excel_file',
            type=str,
            help='Διαδρομή του Excel αρχείου (π.χ. clients.xlsx)'
        )
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Test run - προεπισκόπηση χωρίς αποθήκευση στη βάση'
        )
        parser.add_argument(
            '--skip-errors',
            action='store_true',
            help='Συνέχισε ακόμα και αν υπάρχουν σφάλματα'
        )
        parser.add_argument(
            '--update-only',
            action='store_true',
            help='Μόνο update υπαρχόντων (δεν δημιουργεί νέους)'
        )
        parser.add_argument(
            '--create-only',
            action='store_true',
            help='Μόνο δημιουργία νέων (δεν ενημερώνει υπάρχοντες)'
        )
        parser.add_argument(
            '--verbose',
            action='store_true',
            help='Εμφάνιση αναλυτικών πληροφοριών'
        )

    def handle(self, *args, **options):
        excel_file = options['excel_file']
        dry_run = options['dry_run']
        skip_errors = options['skip_errors']
        update_only = options['update_only']
        create_only = options['create_only']
        verbose = options['verbose']
        
        # Banner
        self.print_banner()
        
        if dry_run:
            self.stdout.write(self.style.WARNING('🔍 DRY RUN MODE - Δεν θα αποθηκευτούν αλλαγές\n'))
        
        if update_only:
            self.stdout.write(self.style.NOTICE('📝 UPDATE ONLY - Μόνο ενημέρωση υπαρχόντων\n'))
        
        if create_only:
            self.stdout.write(self.style.NOTICE('➕ CREATE ONLY - Μόνο δημιουργία νέων\n'))
        
        # Άνοιγμα Excel
        self.stdout.write(f'📂 Άνοιγμα αρχείου: {excel_file}')
        
        try:
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active
        except FileNotFoundError:
            self.stdout.write(self.style.ERROR(f'❌ Το αρχείο δεν βρέθηκε: {excel_file}'))
            return
        except Exception as e:
            self.stdout.write(self.style.ERROR(f'❌ Σφάλμα ανάγνωσης: {e}'))
            return
        
        # Field mapping - ΟΛΑ τα πεδία
        field_mapping = self.get_field_mapping()
        
        # Διάβασμα headers
        headers = self.read_headers(ws, field_mapping)
        
        if not headers:
            self.stdout.write(self.style.ERROR('❌ Δεν βρέθηκαν έγκυρα headers!'))
            return
        
        self.stdout.write(self.style.SUCCESS(f'✅ Βρέθηκαν {len(headers)} έγκυρες στήλες\n'))
        
        if verbose:
            self.stdout.write('📋 Στήλες που θα εισαχθούν:')
            for col_idx, field_name in sorted(headers.items()):
                self.stdout.write(f'  • Column {col_idx}: {field_name}')
            self.stdout.write('')
        
        # Counters
        stats = {
            'created': 0,
            'updated': 0,
            'skipped': 0,
            'errors': 0,
            'total_rows': 0
        }
        errors = []
        
        # Προσδιορισμός start row (παράλειψη παραδειγμάτων)
        start_row = self.find_start_row(ws)
        total_rows = ws.max_row - start_row + 1
        
        self.stdout.write(f'📊 Επεξεργασία {total_rows} γραμμών...\n')
        self.stdout.write('='*70 + '\n')
        
        # Process rows
        for row_num in range(start_row, ws.max_row + 1):
            stats['total_rows'] += 1
            
            # Progress indicator
            if stats['total_rows'] % 10 == 0:
                self.print_progress(stats['total_rows'], total_rows)
            
            try:
                # Parse row
                row_data = self.parse_row(ws, row_num, headers, verbose)
                
                if not row_data:
                    stats['skipped'] += 1
                    if verbose:
                        self.stdout.write(f'  ⏭️  Γραμμή {row_num}: Παράλειψη (κενή γραμμή)')
                    continue
                
                # Validate
                validation_errors = self.validate_row_data(row_data, row_num)
                if validation_errors:
                    stats['errors'] += 1
                    for error in validation_errors:
                        errors.append(f'Γραμμή {row_num}: {error}')
                        self.stdout.write(self.style.ERROR(f'  ❌ {error}'))
                    if not skip_errors:
                        break
                    continue
                
                # Dry run - just show what would happen
                if dry_run:
                    exists = ClientProfile.objects.filter(afm=row_data['afm']).exists()
                    if exists:
                        stats['updated'] += 1
                        self.stdout.write(
                            self.style.WARNING(f'  🔄 [DRY RUN] Θα ενημερωθεί: {row_data["afm"]} - {row_data["eponimia"]}')
                        )
                    else:
                        stats['created'] += 1
                        self.stdout.write(
                            self.style.SUCCESS(f'  ✅ [DRY RUN] Θα δημιουργηθεί: {row_data["afm"]} - {row_data["eponimia"]}')
                        )
                    continue
                
                # Real import with transaction
                with transaction.atomic():
                    afm = row_data.pop('afm')

                    # Τα credential πεδία πάνε στο ClientCredential (κρυπτογραφημένα)
                    from accounting.services.credentials import (
                        extract_legacy_credentials, store_client_credentials,
                    )
                    credentials = extract_legacy_credentials(row_data)
                    
                    # Check if exists
                    exists = ClientProfile.objects.filter(afm=afm).exists()
                    
                    # Apply create/update filters
                    if update_only and not exists:
                        stats['skipped'] += 1
                        if verbose:
                            self.stdout.write(f'  ⏭️  {afm}: Παράλειψη (δεν υπάρχει - update only mode)')
                        continue
                    
                    if create_only and exists:
                        stats['skipped'] += 1
                        if verbose:
                            self.stdout.write(f'  ⏭️  {afm}: Παράλειψη (υπάρχει ήδη - create only mode)')
                        continue
                    
                    # Create or update
                    client, created = ClientProfile.objects.update_or_create(
                        afm=afm,
                        defaults=row_data
                    )

                    if credentials:
                        store_client_credentials(client, credentials)


                    if created:
                        stats['created'] += 1
                        self.stdout.write(
                            self.style.SUCCESS(f'  ✅ Νέος: {client.afm} - {client.eponimia}')
                        )
                    else:
                        stats['updated'] += 1
                        self.stdout.write(
                            self.style.WARNING(f'  🔄 Ενημέρωση: {client.afm} - {client.eponimia}')
                        )
            
            except Exception as e:
                stats['errors'] += 1
                error_msg = f'Γραμμή {row_num}: {str(e)}'
                errors.append(error_msg)
                self.stdout.write(self.style.ERROR(f'  ❌ {error_msg}'))
                
                if not skip_errors:
                    self.stdout.write(self.style.ERROR(f'\n⛔ Διακόπηκε λόγω σφάλματος. Χρησιμοποίησε --skip-errors για συνέχεια.'))
                    break
        
        # Final summary
        self.print_summary(stats, errors, dry_run)
    
    def get_field_mapping(self):
        """Πλήρες mapping Excel headers -> Model fields"""
        return {
            'Α.Φ.Μ.': 'afm',
            'Δ.Ο.Υ.': 'doy',
            'Επωνυμία/Επώνυμο': 'eponimia',
            'Όνομα': 'onoma',
            'Όνομα Πατρός': 'onoma_patros',
            'Αριθμός Ταυτότητας': 'arithmos_taftotitas',
            'Είδος Ταυτότητας': 'eidos_taftotitas',
            'Προσωπικός Αριθμός': 'prosopikos_arithmos',
            'Α.Μ.Κ.Α.': 'amka',
            'Α.Μ. Ι.Κ.Α.': 'am_ika',
            'Αριθμός Γ.Ε.ΜΗ.': 'arithmos_gemi',
            'Αριθμός Δ.ΥΠ.Α': 'arithmos_dypa',
            'Ημ. Γέννησης': 'imerominia_gennisis',
            'Ημ. Γάμου': 'imerominia_gamou',
            'Φύλο': 'filo',
            'Διεύθυνση Κατοικίας': 'diefthinsi_katoikias',
            'Αριθμός': 'arithmos_katoikias',
            'Πόλη Κατοικίας': 'poli_katoikias',
            'Δήμος Κατοικίας': 'dimos_katoikias',
            'Νομός Κατοικίας': 'nomos_katoikias',
            'T.K. Κατοικίας': 'tk_katoikias',
            'Τηλέφωνο Οικίας 1': 'tilefono_oikias_1',
            'Τηλέφωνο Οικίας 2': 'tilefono_oikias_2',
            'Κινητό τηλέφωνο': 'kinito_tilefono',
            'Διεύθυνση Επιχείρησης': 'diefthinsi_epixeirisis',
            'Αριθμός Επιχείρησης': 'arithmos_epixeirisis',
            'Πόλη Επιχείρησης': 'poli_epixeirisis',
            'Δήμος Επιχείρησης': 'dimos_epixeirisis',
            'Νομός Επιχείρησης': 'nomos_epixeirisis',
            'Τ.Κ. Επιχείρησης': 'tk_epixeirisis',
            'Τηλέφωνο Επιχείρησης 1': 'tilefono_epixeirisis_1',
            'Τηλέφωνο Επιχείρησης 2': 'tilefono_epixeirisis_2',
            'Email': 'email',
            'Τράπεζα': 'trapeza',
            'IBAN': 'iban',
            'Είδος Υπόχρεου': 'eidos_ipoxreou',
            'Κατηγορία Βιβλίων': 'katigoria_vivlion',
            'Νομική Μορφή': 'nomiki_morfi',
            'Αγρότης': 'agrotis',
            'Ημ/νία Έναρξης Εργασιών': 'imerominia_enarksis',
            'Όνομα Χρήστη Taxis Net': 'onoma_xristi_taxisnet',
            'Κωδικός Taxis Net': 'kodikos_taxisnet',
            'Όνομα Χρήστη Ι.Κ.Α. Εργοδότη': 'onoma_xristi_ika_ergodoti',
            'Κωδικός Ι.Κ.Α. Εργοδότη': 'kodikos_ika_ergodoti',
            'Όνομα Χρήστη Γ.Ε.ΜΗ.': 'onoma_xristi_gemi',
            'Κωδικός Γ.Ε.ΜΗ.': 'kodikos_gemi',
            'Α.Φ.Μ Συζύγου/Μ.Σ.Σ.': 'afm_sizigou',
            'Α.Φ.Μ. Φορέας': 'afm_foreas',
            'ΑΜ ΚΛΕΙΔΙ': 'am_klidi',
        }
    
    def read_headers(self, ws, field_mapping):
        """Διαβάζει τα headers από την 1η γραμμή"""
        headers = {}
        for col_idx, cell in enumerate(ws[1], start=1):
            header_name = str(cell.value).strip() if cell.value else None
            if header_name and header_name in field_mapping:
                headers[col_idx] = field_mapping[header_name]
        return headers
    
    def find_start_row(self, ws):
        """Βρίσκει από ποια γραμμή να ξεκινήσει (παράλειψη παραδειγμάτων)"""
        # Έλεγχος αν η 2η γραμμή είναι παράδειγμα
        if ws.max_row >= 2:
            first_afm = ws.cell(2, 1).value
            if first_afm and str(first_afm).startswith('123456'):
                self.stdout.write(self.style.WARNING('⚠️  Παράλειψη παράδειγμα γραμμής'))
                return 3
        return 2
    
    def parse_row(self, ws, row_num, headers, verbose):
        """Parse μια γραμμή Excel σε dictionary"""
        row_data = {}
        has_data = False
        
        for col_idx, field_name in headers.items():
            cell = ws.cell(row_num, col_idx)
            value = cell.value
            
            # Skip αν όλα τα κελιά είναι κενά
            if value is not None and str(value).strip():
                has_data = True
            
            # Clean & convert value
            value = self.clean_value(value, field_name, row_num)
            
            if value is not None:
                row_data[field_name] = value
        
        return row_data if has_data else None
    
    def clean_value(self, value, field_name, row_num):
        """Καθαρισμός και μετατροπή τιμών"""
        
        # Null/Empty check
        if value is None:
            return None
        
        # String cleaning
        if isinstance(value, str):
            value = value.strip()
            if value == '' or value.upper() in ['ΚΕΝΟ', 'EMPTY', '-', 'N/A']:
                return None
        
        # Είδος Υπόχρεου - REQUIRED
        if field_name == 'eidos_ipoxreou':
            mapping = {
                'ΙΔΙΩΤΗΣ': 'individual',
                'ΕΠΑΓΓΕΛΜΑΤΙΑΣ': 'professional',
                'ΕΤΑΙΡΕΙΑ': 'company',
                'INDIVIDUAL': 'individual',
                'PROFESSIONAL': 'professional',
                'COMPANY': 'company',
            }
            value_upper = str(value).strip().upper()
            return mapping.get(value_upper, 'professional')  # Default
        
        # Κατηγορία Βιβλίων
        if field_name == 'katigoria_vivlion':
            mapping = {
                'Α': 'A',
                'Β': 'B',
                'Γ': 'C',
                'ΧΩΡΙΣ': 'none',
                'A': 'A',
                'B': 'B',
                'C': 'C',
                'NONE': 'none',
            }
            if value:
                value_clean = str(value).strip().upper()
                return mapping.get(value_clean, '')
            return ''
        
        # Boolean fields
        if field_name == 'agrotis':
            if isinstance(value, bool):
                return value
            value_upper = str(value).upper()
            return value_upper in ['ΝΑΙ', 'NAI', 'YES', 'TRUE', '1', 'Ν']
        
        # Φύλο
        if field_name == 'filo':
            if value:
                value_upper = str(value).upper()
                if value_upper in ['Μ', 'M', 'ΑΝΔΡΑΣ', 'MALE', 'MAN']:
                    return 'M'
                elif value_upper in ['Γ', 'F', 'ΓΥΝΑΙΚΑ', 'FEMALE', 'WOMAN']:
                    return 'F'
            return ''
        
        # Dates
        if field_name in ('imerominia_gennisis', 'imerominia_gamou', 'imerominia_enarksis'):
            return self.parse_date(value, row_num, field_name)
        
        # Default: string
        return str(value) if value else ''
    
    def parse_date(self, value, row_num, field_name):
        """Parse ημερομηνία από διάφορες μορφές"""
        if not value:
            return None
        
        # Αν είναι ήδη datetime object
        if isinstance(value, datetime):
            return value.date()
        
        # Try διάφορες μορφές
        date_formats = [
            '%d/%m/%Y',
            '%d-%m-%Y',
            '%Y-%m-%d',
            '%d.%m.%Y',
            '%d/%m/%y',
        ]
        
        value_str = str(value).strip()
        
        for fmt in date_formats:
            try:
                return datetime.strptime(value_str, fmt).date()
            except ValueError:
                continue
        
        # Αν δεν μπόρεσε να parse, επέστρεψε None
        return None
    
    def validate_row_data(self, row_data, row_num):
        """Validation rules"""
        errors = []
        
        # Required: AFM
        if not row_data.get('afm'):
            errors.append('ΑΦΜ είναι υποχρεωτικό')
        else:
            afm = str(row_data['afm']).strip()
            if len(afm) != 9 or not afm.isdigit():
                errors.append(f'ΑΦΜ πρέπει να είναι 9 ψηφία (δόθηκε: {afm})')
        
        # Required: Επωνυμία
        if not row_data.get('eponimia'):
            errors.append('Επωνυμία είναι υποχρεωτική')
        
        # Required: Είδος Υπόχρεου
        if not row_data.get('eidos_ipoxreou'):
            errors.append('Είδος Υπόχρεου είναι υποχρεωτικό')
        elif row_data['eidos_ipoxreou'] not in ['individual', 'professional', 'company']:
            errors.append(f'Άκυρο Είδος Υπόχρεου: {row_data["eidos_ipoxreou"]}')
        
        # Email validation
        if row_data.get('email'):
            email = row_data['email']
            if '@' not in email or '.' not in email:
                errors.append(f'Άκυρο email: {email}')
        
        # IBAN validation (basic)
        if row_data.get('iban'):
            iban = str(row_data['iban']).replace(' ', '')
            if not iban.startswith('GR') or len(iban) != 27:
                errors.append(f'Άκυρο IBAN: {row_data["iban"]} (πρέπει να ξεκινά με GR και να έχει 27 χαρακτήρες)')
        
        return errors
    
    def print_progress(self, current, total):
        """Progress bar"""
        percent = int((current / total) * 100)
        bar_length = 40
        filled = int(bar_length * current / total)
        bar = '█' * filled + '░' * (bar_length - filled)
        sys.stdout.write(f'\r  [{bar}] {percent}% ({current}/{total})')
        sys.stdout.flush()
        if current == total:
            sys.stdout.write('\n')
    
    def print_banner(self):
        """ASCII banner"""
        banner = """
╔═══════════════════════════════════════════════════════════════════╗
║                    📥 CLIENT IMPORT TOOL                          ║
║                   LogistikoCRM - v2.0                             ║
╚═══════════════════════════════════════════════════════════════════╝
        """
        self.stdout.write(self.style.SUCCESS(banner))
    
    def print_summary(self, stats, errors, dry_run):
        """Τελική αναφορά"""
        self.stdout.write('\n' + '='*70)
        
        if dry_run:
            self.stdout.write(self.style.SUCCESS('\n✅ DRY RUN ΟΛΟΚΛΗΡΩΘΗΚΕ'))
        else:
            self.stdout.write(self.style.SUCCESS('\n✅ IMPORT ΟΛΟΚΛΗΡΩΘΗΚΕ!'))
        
        self.stdout.write('\n📊 ΑΠΟΤΕΛΕΣΜΑΤΑ:')
        self.stdout.write('─' * 70)
        self.stdout.write(f'  📝 Σύνολο γραμμών:        {stats["total_rows"]}')
        self.stdout.write(f'  ✅ Νέοι πελάτες:          {stats["created"]}')
        self.stdout.write(f'  🔄 Ενημερωμένοι:          {stats["updated"]}')
        self.stdout.write(f'  ⏭️  Παραλείφθηκαν:        {stats["skipped"]}')
        self.stdout.write(f'  ❌ Σφάλματα:              {stats["errors"]}')
        
        success_rate = 0
        if stats['total_rows'] > 0:
            successful = stats['created'] + stats['updated']
            success_rate = (successful / stats['total_rows']) * 100
        
        self.stdout.write(f'  📈 Επιτυχία:              {success_rate:.1f}%')
        
        # Errors list
        if errors:
            self.stdout.write('\n❌ ΛΙΣΤΑ ΣΦΑΛΜΑΤΩΝ:')
            self.stdout.write('─' * 70)
            for error in errors[:20]:  # Show max 20 errors
                self.stdout.write(f'  • {error}')
            
            if len(errors) > 20:
                self.stdout.write(f'  ... και {len(errors) - 20} ακόμα σφάλματα')
        
        self.stdout.write('\n' + '='*70)
        
        if dry_run:
            self.stdout.write(self.style.NOTICE('\n💡 Αυτό ήταν DRY RUN. Τρέξε χωρίς --dry-run για πραγματικό import.\n'))
        elif stats['errors'] == 0:
            self.stdout.write(self.style.SUCCESS('\n🎉 Το import ολοκληρώθηκε χωρίς σφάλματα!\n'))
        else:
            self.stdout.write(self.style.WARNING(f'\n⚠️  Το import ολοκληρώθηκε με {stats["errors"]} σφάλματα.\n'))